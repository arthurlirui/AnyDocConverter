"""
PDF → XLSX 转换器
使用 PyMuPDF 表格检测 + openpyxl 写入
"""

from __future__ import annotations

import logging
import os
from typing import Optional

import fitz  # PyMuPDF

from algo.models.params import ConvertParams

logger = logging.getLogger(__name__)


def convert(file_path: str, params: ConvertParams, output_dir: str) -> str:
    """PDF → XLSX 转换

    策略:
    1. 用 PyMuPDF 的 find_tables 检测表格
    2. 每个表格写入一个 sheet
    3. 如果多页有表格，按页名命名 sheet

    Returns:
        输出 .xlsx 文件路径
    """
    os.makedirs(output_dir, exist_ok=True)
    base_name = os.path.splitext(os.path.basename(file_path))[0]
    output_path = os.path.join(output_dir, f"{base_name}.xlsx")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        logger.error("openpyxl not installed. pip install openpyxl")
        raise

    doc = fitz.open(file_path)
    wb = Workbook()
    # 删除默认 sheet
    wb.remove(wb.active)

    start = params.start_page
    end = params.end_page if params.end_page is not None else doc.page_count - 1
    sheet_count = 0

    for i in range(start, end + 1):
        page = doc[i]

        try:
            tables = page.find_tables()
        except Exception as e:
            logger.debug("Page %d: table detection failed: %s", i, e)
            continue

        # 物化表格列表，避免多次调用 find_tables()
        table_list = list(tables)

        for table in table_list:
            sheet_count += 1
            sheet_name = f"Page{i + 1}_Table{sheet_count}"[:31]  # Excel 限制 31 字符
            ws = wb.create_sheet(title=sheet_name)

            data = table.extract()
            for row_idx, row_data in enumerate(data, start=1):
                for col_idx, cell_value in enumerate(row_data, start=1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)

                    # 表头加粗
                    if row_idx == 1:
                        cell.font = Font(bold=True)

        # 如果没有找到表格，尝试提取纯文本
        if not table_list:
            text = page.get_text("text").strip()
            if text:
                sheet_count += 1
                sheet_name = f"Page{i + 1}_Text"[:31]
                ws = wb.create_sheet(title=sheet_name)
                for line_idx, line in enumerate(text.split("\n"), start=1):
                    ws.cell(row=line_idx, column=1, value=line)

    doc.close()

    if sheet_count == 0:
        # 没有任何内容，创建一张空白 sheet
        ws = wb.create_sheet(title="NoTables")
        ws.cell(row=1, column=1, value="No tables or text found in the selected pages.")

    wb.save(output_path)
    logger.info("XLSX conversion done: %s", output_path)
    return output_path
