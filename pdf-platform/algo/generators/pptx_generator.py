"""
PPTXGenerator — DocumentModel → 可编辑 .pptx。

从 :class:`~algo.document_model.DocumentModel` 生成幻灯片：

- 每页 → 一张幻灯片（尺寸按源页面比例）
- text 块 → 定位文本框（保留坐标、字号、颜色）
- title 块 → 定位标题文本框
- table 块 → 定位表格形状（保留行列结构与合并）
- image 块 → 定位图片（按原始坐标放置）
- 公式 → 文本框中显示 LaTeX 纯文本

设计要点
--------
PPTX 是固定布局格式，所有元素都有绝对坐标，天然适合像素级版式还原。
归一化坐标 [0,1] × 幻灯片尺寸 = 绝对 EMU 坐标。
文本框保留字号/颜色/粗体，用户可直接点击编辑。
"""

from __future__ import annotations

import io
import logging
import os
from typing import Optional

from algo.document_model import DocumentModel, LayoutBlock, Page, TableBlock, TextRun
from algo.models.params import ConvertParams

logger = logging.getLogger(__name__)

_EMU_PER_PT = 914400 / 72  # 1 pt = 12700 EMU


def generate_pptx(
    model: DocumentModel,
    output_path: str,
    params: Optional[ConvertParams] = None,
) -> str:
    """从 DocumentModel 生成 PPTX 文件。

    Args:
        model: 解析阶段产出的文档模型。
        output_path: 输出 .pptx 文件路径。
        params: 转换参数。

    Returns:
        输出文件路径。
    """
    if params is None:
        params = ConvertParams()

    from pptx import Presentation
    from pptx.util import Emu, Pt

    prs = Presentation()

    for page_idx, page in enumerate(model.pages):
        # 设置幻灯片尺寸（只在首页设置，后续页统一）
        if page_idx == 0:
            prs.slide_width = Emu(int(page.width * _EMU_PER_PT))
            prs.slide_height = Emu(int(page.height * _EMU_PER_PT))

        slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

        for block in page.sorted_blocks:
            _write_block_to_pptx(slide, block, page, prs, params)

    prs.save(output_path)
    logger.info("PPTX generated: %s", output_path)
    return output_path


def _write_block_to_pptx(
    slide,
    block: LayoutBlock,
    page: Page,
    prs,
    params: ConvertParams,
):
    """将单个版面块写入幻灯片。"""
    if block.block_type == "table" and block.table:
        _write_table(slide, block.table, page, params)
    elif block.block_type == "image" and block.image:
        _write_image(slide, block, page, params)
    elif block.block_type in ("text", "title", "list", "header", "footer", "caption"):
        _write_text_box(slide, block, page, params)
    elif block.block_type == "formula" and block.formula:
        _write_formula(slide, block, page, params)


def _write_text_box(
    slide,
    block: LayoutBlock,
    page: Page,
    params: ConvertParams,
):
    """写入定位文本框。"""
    from pptx.util import Emu, Pt
    from pptx.dml.color import RGBColor

    text = block.text.strip()
    if not text:
        return

    # 归一化坐标 → EMU
    left = Emu(int(block.bbox.x0 * page.width * _EMU_PER_PT))
    top = Emu(int(block.bbox.y0 * page.height * _EMU_PER_PT))
    width = Emu(int(max(block.bbox.width, 0.02) * page.width * _EMU_PER_PT))
    # 高度略增以容纳文字
    height = Emu(int((block.bbox.height + 0.01) * page.height * _EMU_PER_PT))

    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True

    # 第一段
    p = tf.paragraphs[0]
    _apply_runs_to_pptx_paragraph(p, block.runs, params)

    # 标题加粗
    if block.block_type == "title":
        for run in p.runs:
            run.font.bold = True


def _write_table(
    slide,
    table: TableBlock,
    page: Page,
    params: ConvertParams,
):
    """写入定位表格。"""
    from pptx.util import Emu, Pt

    if table.rows == 0 or table.cols == 0:
        return

    left = Emu(int(table.bbox.x0 * page.width * _EMU_PER_PT))
    top = Emu(int(table.bbox.y0 * page.height * _EMU_PER_PT))
    width = Emu(int(max(table.bbox.width, 0.1) * page.width * _EMU_PER_PT))
    height = Emu(int(max(table.bbox.height, 0.05) * page.height * _EMU_PER_PT))

    try:
        table_shape = slide.shapes.add_table(
            table.rows, table.cols, left, top, width, height
        )
        pptx_table = table_shape.table

        # 标记被合并占据的格子
        occupied = [[False] * table.cols for _ in range(table.rows)]

        for cell in table.cells:
            r, c = cell.row, cell.col
            if r >= table.rows or c >= table.cols:
                continue
            if occupied[r][c]:
                continue

            pptx_cell = pptx_table.cell(r, c)

            # 合并
            if cell.rowspan > 1 or cell.colspan > 1:
                end_r = min(r + cell.rowspan - 1, table.rows - 1)
                end_c = min(c + cell.colspan - 1, table.cols - 1)
                pptx_cell = pptx_cell.merge(pptx_table.cell(end_r, end_c))
                for rr in range(r, end_r + 1):
                    for cc in range(c, end_c + 1):
                        if rr != r or cc != c:
                            occupied[rr][cc] = True

            # 写入文本
            cell_text = "".join(run.text for run in cell.runs).strip()
            if cell_text:
                tf = pptx_cell.text_frame
                p = tf.paragraphs[0]
                _apply_runs_to_pptx_paragraph(p, cell.runs, params)

            # 表头加粗
            if cell.is_header:
                for para in pptx_cell.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.bold = True

    except Exception as e:
        logger.warning("Failed to create PPTX table: %s", e)


def _write_image(
    slide,
    block: LayoutBlock,
    page: Page,
    params: ConvertParams,
):
    """写入定位图片。"""
    from pptx.util import Emu

    img = block.image
    if img is None:
        return

    # 获取图片字节
    if img.image_bytes:
        img_stream = io.BytesIO(img.image_bytes)
    elif img.image_path and os.path.isfile(img.image_path):
        with open(img.image_path, "rb") as f:
            img_stream = io.BytesIO(f.read())
    else:
        return

    left = Emu(int(block.bbox.x0 * page.width * _EMU_PER_PT))
    top = Emu(int(block.bbox.y0 * page.height * _EMU_PER_PT))
    width = Emu(int(max(block.bbox.width, 0.05) * page.width * _EMU_PER_PT))
    height = Emu(int(max(block.bbox.height, 0.05) * page.height * _EMU_PER_PT))

    try:
        slide.shapes.add_picture(img_stream, left, top, width, height)
    except Exception as e:
        logger.warning("Failed to insert PPTX image: %s", e)


def _write_formula(
    slide,
    block: LayoutBlock,
    page: Page,
    params: ConvertParams,
):
    """写入公式文本框。"""
    from pptx.util import Emu, Pt

    formula = block.formula
    if formula is None or not formula.latex:
        return

    left = Emu(int(block.bbox.x0 * page.width * _EMU_PER_PT))
    top = Emu(int(block.bbox.y0 * page.height * _EMU_PER_PT))
    width = Emu(int(max(block.bbox.width, 0.1) * page.width * _EMU_PER_PT))
    height = Emu(int(max(block.bbox.height, 0.02) * page.height * _EMU_PER_PT))

    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = formula.latex
    run.font.italic = True


def _apply_runs_to_pptx_paragraph(p, runs: list[TextRun], params: ConvertParams):
    """将 TextRun 列表应用到 pptx 段落。"""
    from pptx.util import Pt
    from pptx.dml.color import RGBColor

    if not runs:
        return

    for run_data in runs:
        if not run_data.text:
            continue
        run = p.add_run()
        run.text = run_data.text

        # 字号
        if run_data.font_size:
            run.font.size = Pt(run_data.font_size)

        # 字体名
        if run_data.font_name:
            run.font.name = run_data.font_name
        else:
            run.font.name = params.font.fallback_font

        # 粗体 / 斜体
        if run_data.bold:
            run.font.bold = True
        if run_data.italic:
            run.font.italic = True

        # 颜色
        if run_data.color:
            rgb = _hex_to_rgb(run_data.color)
            if rgb:
                run.font.color.rgb = RGBColor(*rgb)


def _hex_to_rgb(hex_str: str) -> Optional[tuple[int, int, int]]:
    """'#RRGGBB' → (r, g, b)。"""
    try:
        hex_str = hex_str.lstrip("#")
        if len(hex_str) == 6:
            return (int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))
    except (ValueError, IndexError):
        pass
    return None
